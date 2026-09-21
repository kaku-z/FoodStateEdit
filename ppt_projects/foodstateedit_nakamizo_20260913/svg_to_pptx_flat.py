#!/usr/bin/env python3
"""Flat-mode SVG -> native DrawingML PPTX converter.

The build_slides.py SVG output intentionally uses only a small shape subset
(rect / line / text / image / circle / polygon, no style/class/foreignObject).
This converter maps each element 1:1 onto native PowerPoint shapes at the
same pixel coordinates (1280x720 px = 13.333x7.5 in slide), the same
"flat" strategy as the project's usual export pipeline.

Usage:
    py svg_to_pptx_flat.py [--out exports/name.pptx]
"""

from __future__ import annotations

import argparse
import base64
import re
import time
from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parent
SVG_DIR = ROOT / "svg_output"
IMG_DIR = ROOT / "images"
EXPORTS = ROOT / "exports"

NS = {"s": "http://www.w3.org/2000/svg"}
PX_EMU = 9525  # 96 dpi
SLIDE_W = 1280
SLIDE_H = 720


def color(v: str) -> RGBColor:
    v = v.strip()
    if v.startswith("#") and len(v) == 7:
        return RGBColor(int(v[1:3], 16), int(v[3:5], 16), int(v[5:7], 16))
    if v == "none":
        raise ValueError("none")
    return RGBColor(0, 0, 0)


def map_font(family: str) -> str:
    if "Yu Gothic" in family or "Meiryo" in family:
        return "Yu Gothic"
    return "Arial"


def add_rect(slide, x, y, w, h, fill, stroke, sw, rx):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rx else MSO_SHAPE.RECTANGLE,
        Emu(x * PX_EMU), Emu(y * PX_EMU), Emu(w * PX_EMU), Emu(h * PX_EMU))
    try:
        shape.fill.solid()
        shape.fill.fore_color.rgb = color(fill)
    except ValueError:
        shape.fill.background()
    if stroke and stroke != "none":
        shape.line.color.rgb = color(stroke)
        shape.line.width = Pt(sw)
    else:
        shape.line.fill.background()
    if rx:
        shape.adjustments[0] = min(0.5, rx / min(w, h))
    shape.shadow.inherit = False
    return shape


def add_line(slide, x1, y1, x2, y2, stroke, sw):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Emu(x1 * PX_EMU), Emu(y1 * PX_EMU), Emu(x2 * PX_EMU), Emu(y2 * PX_EMU))
    conn.line.color.rgb = color(stroke)
    conn.line.width = Pt(sw)
    conn.shadow.inherit = False
    return conn


def add_text(slide, x, y, text, size, fill, weight, anchor, family):
    w = Emu(max(int(len(text) * size * 1.25 * PX_EMU) + 20000, 40000))
    h = Emu(int(size * 1.6 * PX_EMU))
    left = Emu(x * PX_EMU)
    if anchor == "middle":
        left -= w // 2
    elif anchor == "end":
        left -= w
    box = slide.shapes.add_textbox(left, Emu(int((y - size * 1.18) * PX_EMU)), w, h)
    tf = box.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = {"start": PP_ALIGN.LEFT, "middle": PP_ALIGN.CENTER, "end": PP_ALIGN.RIGHT}[anchor]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = weight >= 600
    run.font.name = map_font(family)
    run.font.color.rgb = color(fill)
    return box


def add_image(slide, href, x, y, w, h):
    if href.startswith("data:"):
        data = base64.b64decode(href.split(",", 1)[1])
        path = ROOT / "_inline_tmp.png"
        path.write_bytes(data)
    else:
        path = IMG_DIR / Path(href).name
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    dw, dh = iw * scale, ih * scale
    slide.shapes.add_picture(str(path), Emu(int((x + (w - dw) / 2) * PX_EMU)),
                             Emu(int((y + (h - dh) / 2) * PX_EMU)),
                             Emu(int(dw * PX_EMU)), Emu(int(dh * PX_EMU)))


def add_circle(slide, cx, cy, r, fill, stroke, sw):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                                   Emu((cx - r) * PX_EMU), Emu((cy - r) * PX_EMU),
                                   Emu(2 * r * PX_EMU), Emu(2 * r * PX_EMU))
    try:
        shape.fill.solid()
        shape.fill.fore_color.rgb = color(fill)
    except ValueError:
        shape.fill.background()
    if stroke and stroke != "none":
        shape.line.color.rgb = color(stroke)
        shape.line.width = Pt(sw)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False


def add_polygon(slide, points, fill):
    builder = slide.shapes.build_freeform(Emu(points[0][0] * PX_EMU), Emu(points[0][1] * PX_EMU))
    for px, py in points[1:]:
        builder.add_line_segments([(Emu(px * PX_EMU), Emu(py * PX_EMU))], close=False)
    shape = builder.convert_to_shape()
    try:
        shape.fill.solid()
        shape.fill.fore_color.rgb = color(fill)
    except ValueError:
        shape.fill.background()
    shape.line.fill.background()
    shape.shadow.inherit = False


def num(v: str) -> float:
    return float(v.strip())


def convert(svg_path: Path, slide) -> None:
    tree = ET.parse(svg_path)
    for el in tree.iter():
        tag = el.tag.split("}")[-1]
        if tag == "rect":
            add_rect(slide, num(el.get("x")), num(el.get("y")), num(el.get("width")),
                     num(el.get("height")), el.get("fill", "none"), el.get("stroke", "none"),
                     num(el.get("stroke-width", "1")), num(el.get("rx", "0")))
        elif tag == "line":
            add_line(slide, num(el.get("x1")), num(el.get("y1")), num(el.get("x2")),
                     num(el.get("y2")), el.get("stroke", "#000000"), num(el.get("stroke-width", "2")))
        elif tag == "text":
            add_text(slide, num(el.get("x")), num(el.get("y")), el.text or "",
                     num(el.get("font-size", "16")), el.get("fill", "#222222"),
                     int(el.get("font-weight", "400")), el.get("text-anchor", "start"),
                     el.get("font-family", "Arial"))
        elif tag == "image":
            add_image(slide, el.get("href"), num(el.get("x")), num(el.get("y")),
                      num(el.get("width")), num(el.get("height")))
        elif tag == "circle":
            add_circle(slide, num(el.get("cx")), num(el.get("cy")), num(el.get("r")),
                       el.get("fill", "none"), el.get("stroke", "none"), num(el.get("stroke-width", "1")))
        elif tag == "polygon":
            pts = [tuple(map(float, p.split(","))) for p in el.get("points", "").split()]
            add_polygon(slide, pts, el.get("fill", "#0070C0"))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, default=None)
    args = parser.parse_args()
    if args.out is None:
        stamp = time.strftime("%Y%m%d_%H%M%S")
        args.out = EXPORTS / f"foodstateedit_nakamizo_20260913_{stamp}.pptx"
    args.out.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W * PX_EMU)
    prs.slide_height = Emu(SLIDE_H * PX_EMU)
    blank = prs.slide_layouts[6]
    svgs = sorted(SVG_DIR.glob("slide_*.svg"))
    for svg_path in svgs:
        slide = prs.slides.add_slide(blank)
        convert(svg_path, slide)
    prs.save(args.out)
    print(f"wrote {args.out} ({len(svgs)} slides)")


if __name__ == "__main__":
    main()
